import re
import docx

def analyze(file_path):
    try:
        doc = docx.Document(file_path)
    except:
        return {"error": "Invalid DOCX file"}

    text = "\n".join([p.text for p in doc.paragraphs])
    
    # 1. Detect Bracketed Citations e.g., [1], [1-3], [4, 5]
    bracket_citations = re.findall(r'\[\s*\d+(?:\s*,\s*\d+)*\s*(?:-\s*\d+)?\s*\]', text)
    
    # 2. Detect Author-Year Citations e.g., (Smith, 2020), (Smith and Jones 2021)
    author_year_citations = re.findall(r'\([A-Za-z\s]+(?:et al\.?)?,?\s*\d{4}\)', text)
    
    total_citations = len(bracket_citations) + len(author_year_citations)
    
    # 3. Detect References List
    references_text = []
    in_references = False
    
    for p in doc.paragraphs:
        p_text = p.text.strip()
        if not p_text:
            continue
            
        # Check if we hit the references section
        if p_text.lower() in ['references', 'bibliography', 'literature cited']:
            in_references = True
            continue
            
        if in_references:
            # Check if paragraph looks like a reference (starts with [1] or is long enough)
            if re.match(r'^\[\d+\]', p_text) or len(p_text) > 40:
                references_text.append(p_text)

    total_references = len(references_text)

    # Calculate issues
    issues = []
    missing = 0
    unused = 0
    
    if total_citations > 0 and total_references == 0:
        issues.append({"type": "danger", "message": "In-text citations found, but no References section detected."})
        missing = total_citations
    elif total_citations == 0 and total_references > 0:
        issues.append({"type": "warning", "message": "References found, but no in-text citations detected."})
        unused = total_references
    else:
        # Simple ratio check
        if total_citations > total_references * 3:
            missing = total_citations // 3  # Rough heuristic
            issues.append({"type": "warning", "message": f"High ratio of citations to references. Some references may be missing."})
        elif total_references > total_citations * 2:
            unused = total_references - total_citations
            issues.append({"type": "warning", "message": f"Many references do not appear to be cited in the text."})
        else:
            issues.append({"type": "info", "message": "Citation and reference counts appear balanced."})

    # Calculate Score
    base_score = 100
    if total_citations == 0 and total_references == 0:
        base_score = 0
        issues.append({"type": "danger", "message": "No citations or references found."})
    else:
        penalty = (missing * 5) + (unused * 2)
        base_score = max(20, 100 - penalty)
        
    return {
        "citation_score": base_score,
        "total_citations": total_citations,
        "total_references": total_references,
        "missing_references": missing,
        "unused_references": unused,
        "duplicates": 0,
        "issues": issues
    }
